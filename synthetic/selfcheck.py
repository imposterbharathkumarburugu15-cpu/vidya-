#!/usr/bin/env python3
"""
Independent check of a generated dataset: re-reads every upload file from disk, recomputes the
expected codes from the raw cells + registries, and compares them with ground_truth.json.

  python -m synthetic.selfcheck data/

Exit code 0 = ground truth is consistent with the files. This is also a reference implementation
of the parsing rules a validator must handle (header variants, date/bool/seat/reason dialects).
"""
from __future__ import annotations

import csv
import json
import math
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import yaml
from openpyxl import load_workbook

try:
    from .generate import (BLOCKING, GENDER_STYLES, PROGRAMMES, REASON_STYLES, SEAT_STYLES,
                           STATUS_STYLES)
    from . import names as N
except ImportError:  # pragma: no cover
    from generate import (BLOCKING, GENDER_STYLES, PROGRAMMES, REASON_STYLES, SEAT_STYLES,  # type: ignore
                          STATUS_STYLES)
    import names as N  # type: ignore

ROW_LOCAL = {"E101", "E101W", "E102", "E103", "E104", "E304", "E401", "E402", "E403", "E404"}
HONORIFICS = {h.lower().rstrip(".") for h in N.HONORIFICS}


def read_batch(root: Path, b: dict) -> list[dict]:
    dia, path = b["dialect"], root / b["file"]
    if path.suffix == ".xlsx":
        ws = load_workbook(path, read_only=True, data_only=True).active
        lines = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding=dia["csv_encoding"]) as fh:
            lines = list(csv.reader(fh, delimiter=dia["csv_delimiter"]))
    header = [str(h) if h is not None else "" for h in lines[b["header_row_index"]]]
    rev = {v: k for k, v in dia["headers"].items()}
    out = []
    for raw in lines[b["header_row_index"] + 1:]:
        if raw is None or all(c in (None, "") for c in raw):
            continue
        out.append({rev[h]: v for h, v in zip(header, raw) if h in rev})
    return out


def parse_date(v, fmt: str):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return datetime.strptime(str(v), "%Y-%m-%d" if fmt == "excel" else fmt).date()


def canon(row: dict, dia: dict) -> dict:
    inv = lambda table, v: {y: x for x, y in table.items()}.get(v)  # noqa: E731
    t, f = dia["bool_tokens"]
    a = row.get("apaar_id")
    a = "" if a in (None, "") else (str(int(a)) if isinstance(a, (int, float)) else str(a))
    prog_v = row.get("programme")
    prog = next((c for c, (n, _, s) in PROGRAMMES.items() if prog_v in (c, n, s)), None)
    return {
        "apaar": a.replace(" ", "").replace("-", ""),
        "reason": inv(REASON_STYLES[dia["reason_style"]], row.get("apaar_reason") or None),
        "name": row.get("student_name") or "",
        "guardian": row.get("guardian_name") or "",
        "dob": parse_date(row.get("dob"), dia["date_format"]),
        "adm": parse_date(row.get("admission_date"), dia["date_format"]),
        "programme": prog,
        "year": int(row.get("year_of_study")),
        "seat": inv(SEAT_STYLES[dia["seat_style"]], row.get("seat_category")),
        **{k: (str(row.get(k) if row.get(k) is not None else "") == t) for k in
           ("is_nri", "is_oci", "is_fn", "is_supernumerary", "is_lateral_entry")},
        "status": inv(STATUS_STYLES[dia["status_style"]], row.get("enrolment_status")),
        "gender": inv(GENDER_STYLES[dia["gender_style"]], row.get("gender")),
        "declared": row.get("declared_intake"),
    }


def norm_tokens(name: str) -> list[str]:
    toks = [t for t in re.sub(r"\s+", " ", name.strip()).lower().split(" ") if t]
    toks = [t for t in toks if t.rstrip(".") not in HONORIFICS and len(t.rstrip(".")) > 1]
    return toks


def same_person_name(a: str, b: str) -> bool:
    ta, tb = norm_tokens(a), norm_tokens(b)
    if len(ta) != len(tb):
        return False
    groups = {n.lower(): {x.lower() for x in g} for n, g in N.VARIANT_INDEX.items()}
    return all(x == y or y in groups.get(x, set()) for x, y in zip(ta, tb))


def main(root: Path) -> int:
    man = json.loads((root / "manifest.json").read_text())
    truth = json.loads((root / "ground_truth.json").read_text())
    reg = {p["apaar_id"]: p for p in json.loads((root / "registry/apaar_registry.json").read_text())}
    aicte = {i["institution_id"]: i for i in json.loads((root / "registry/aicte_institutions.json").read_text())}
    caps_cfg = yaml.safe_load((root / "registry/seat_categories_used.yaml").read_text())
    year = man["academic_year"]
    w0, w1 = (date.fromisoformat(d) for d in man["admission_window"])
    lo, hi = man["age_bounds"]
    errors: list[str] = []

    rows: dict[tuple[str, int], dict] = {}
    for b in man["batches"]:
        for k, raw in enumerate(read_batch(root, b), start=1):
            r = canon(raw, b["dialect"])
            r.update(file=b["file"], row=k, inst=b["institution_id"], upload_at=b["upload_at"])
            rows[(b["file"], k)] = r
        if b["rows"] != k:
            errors.append(f"{b['file']}: manifest says {b['rows']} rows, file has {k}")

    tmap = {(t["file"], t["row"]): set(t["codes"]) for t in truth["rows"]}

    # 1) row-local codes recomputed from raw cells
    for key, r in rows.items():
        exp: set[str] = set()
        if aicte[r["inst"]]["status_by_year"][year] != "approved":
            exp.add("E304")
        if not r["apaar"]:
            exp.add("E101W" if r["reason"] else "E101")
        elif not re.fullmatch(r"\d{12}", r["apaar"]):
            exp.add("E102")
        elif r["apaar"] not in reg:
            exp.add("E103")
        else:
            p = reg[r["apaar"]]
            if date.fromisoformat(p["dob"]) != r["dob"] or not same_person_name(r["name"], p["name"]):
                exp.add("E104")
        if (r["is_fn"] and r["is_oci"]) or (r["is_nri"] and r["is_fn"]) or (r["seat"] == "FN" and not r["is_fn"]):
            exp.add("E401")
        age = (r["adm"] - r["dob"]).days / 365.25
        if age < lo or age > hi:
            exp.add("E402")
        if not (w0 <= r["adm"] <= w1):
            exp.add("E403")
        if r["seat"] == "LATERAL" and r["year"] == 1:
            exp.add("E404")
        got = tmap.get(key, set()) & ROW_LOCAL
        if exp != got:
            errors.append(f"{key}: recomputed {sorted(exp)} vs truth {sorted(got)}")

    # 2) duplicate groups are real, and no APAAR is reused outside a group/negative
    explained: set[tuple] = set()
    for g in truth["duplicate_groups"]:
        m = [rows[(x["file"], x["row"])] for x in g["members"]]
        a, b = m
        ok = {
            "in_file": a["file"] == b["file"] and a["apaar"] == b["apaar"],
            "in_institution": a["inst"] == b["inst"] and a["file"] != b["file"] and a["apaar"] == b["apaar"],
            "exact": a["inst"] != b["inst"] and a["apaar"] == b["apaar"] and a["status"] == b["status"] == "active",
            "fuzzy": a["inst"] != b["inst"] and (not a["apaar"] or not b["apaar"]) and a["dob"] == b["dob"]
                     and norm_tokens(a["guardian"]) == norm_tokens(b["guardian"]),
        }[g["kind"]]
        if not ok:
            errors.append(f"group {g['kind']} not supported by data: {g['members']}")
        fm = (g["flagged_member"]["file"], g["flagged_member"]["row"])
        if g["code"] not in tmap.get(fm, set()):
            errors.append(f"group {g['kind']}: flagged member {fm} lacks {g['code']}")
        explained.update((x["file"], x["row"]) for x in g["members"])
    for n in truth["negatives"]:
        explained.update((x["file"], x["row"]) for x in n["members"])
    by_apaar: dict[str, list[dict]] = defaultdict(list)
    for r in rows.values():
        if re.fullmatch(r"\d{12}", r["apaar"] or ""):
            by_apaar[r["apaar"]].append(r)
    for a, rs in by_apaar.items():
        if len(rs) > 1 and not all((r["file"], r["row"]) in explained for r in rs):
            errors.append(f"APAAR {a} reused outside any group/negative: {[(r['file'], r['row']) for r in rs]}")

    # 3) programme-level codes recomputed from counted rows
    cats, pools = caps_cfg["seat_categories"], caps_cfg["pools"]
    counts: dict[tuple[str, str], Counter] = defaultdict(Counter)
    for key, r in rows.items():
        if r["status"] != "active" or (tmap.get(key, set()) & BLOCKING):
            continue
        seat = r["seat"]
        pool = "FOREIGN" if seat in ("FN", "OCI", "CIWG", "NRI") else seat
        counts[(r["inst"], r["programme"])][pool] += 1
        if seat == "CIWG":
            counts[(r["inst"], r["programme"])]["CIWG"] += 1
    derived = set()
    for (iid, pc), c in counts.items():
        if aicte[iid]["status_by_year"][year] != "approved":
            continue
        I = next(p["approved_intake"][year] for p in aicte[iid]["programmes"] if p["code"] == pc)
        F = math.floor(I * pools["FOREIGN"]["pct_of_intake"] / 100)
        cap = {"REGULAR": I, "TFW": math.floor(I * cats["TFW"]["pct_of_intake"] / 100), "FOREIGN": F,
               "CIWG": math.floor(F * pools["FOREIGN"]["sub_caps"]["CIWG"]),
               "LATERAL": math.floor(I * cats["LATERAL"]["pct_of_intake"] / 100)}
        for pool, n in c.items():
            if n > cap[pool]:
                derived.add(("E301" if pool == "REGULAR" else "E302", iid, pc, pool))
    declared = {(r["inst"], r["programme"], r["declared"]) for r in rows.values() if r["declared"] not in (None, "")}
    for iid, pc, d in declared:
        I = next(p["approved_intake"][year] for p in aicte[iid]["programmes"] if p["code"] == pc)
        if int(d) != I:
            derived.add(("E303", iid, pc, None))
    truth_p = {(f["code"], f["institution_id"], f["programme"], f.get("pool") if f["code"] != "E303" else None)
               for f in truth["programme_faults"]}
    if derived != truth_p:
        errors.append(f"programme faults differ: only-recomputed={sorted(derived - truth_p)} "
                      f"only-truth={sorted(truth_p - derived)}")

    # 4) mutations: E501 exactly for post-deadline ones
    dl = datetime.fromisoformat(man["reporting_deadline"])
    for m in json.loads((root / "mutations.json").read_text()):
        if (datetime.fromisoformat(m["at"]) > dl) != bool(m["expected_codes"]):
            errors.append(f"mutation timing/expectation mismatch: {m}")

    n_codes = sum(1 for t in truth["rows"] if t["codes"])
    print(f"checked {len(rows)} rows in {len(man['batches'])} batches; {n_codes} rows carry codes; "
          f"{len(truth['duplicate_groups'])} groups; {len(truth['programme_faults'])} programme faults")
    if errors:
        print(f"FAILED with {len(errors)} problem(s):")
        for e in errors[:50]:
            print("  -", e)
        return 1
    print("OK: ground truth matches the files")
    return 0


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[1])))
